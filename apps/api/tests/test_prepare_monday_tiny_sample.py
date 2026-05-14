from __future__ import annotations

import socket
import sqlite3
import csv
from collections import Counter
from pathlib import Path

import pytest

from scripts.prepare_monday_tiny_sample import (
    CLIENT_STATUS_REVIEW_COLUMNS,
    DUPLICATE_SITES_REVIEW_COLUMNS,
    REVIEW_WORKBOOK_SHEETS,
    STATUS_DEFAULTS_REVIEW_COLUMNS,
    UNRESOLVED_SITES_REVIEW_COLUMNS,
    generate_monday_review_pack,
    generate_monday_review_workbook,
    prepare_reviewed_sample,
    prepare_tiny_sample,
)
from scripts.validate_import_templates import validate_import_templates


REPO_ROOT = Path(__file__).resolve().parents[3]
TEMPLATE_DIR = REPO_ROOT / "docs" / "import_templates" / "v2"


def _write_monday_csv(path: Path, title: str, headers: list[str], rows: list[list[str]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerow([title])
        writer.writerow(headers)
        writer.writerows(rows)


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        return list(csv.DictReader(file))


def _write_review_csv(path: Path, columns: tuple[str, ...], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row.get(column, "") for column in columns})


def _review_row(columns: tuple[str, ...], **values: str) -> dict[str, str]:
    row = {column: "" for column in columns}
    row.update(values)
    return row


def _write_review_pack(
    review_dir: Path,
    *,
    unresolved_rows: list[dict[str, str]] | None = None,
    client_status_rows: list[dict[str, str]] | None = None,
    duplicate_rows: list[dict[str, str]] | None = None,
    status_default_rows: list[dict[str, str]] | None = None,
) -> Path:
    _write_review_csv(
        review_dir / "unresolved_sites_review.csv",
        UNRESOLVED_SITES_REVIEW_COLUMNS,
        unresolved_rows
        if unresolved_rows is not None
        else [
            _review_row(
                UNRESOLVED_SITES_REVIEW_COLUMNS,
                source_row_number="5",
                site_id="site-no-client",
                site_name_or_label="No Client Site",
                exclusion_reason="site_has_no_client_id",
            ),
            _review_row(
                UNRESOLVED_SITES_REVIEW_COLUMNS,
                source_row_number="6",
                site_id="site-unmapped-client",
                site_name_or_label="Unmapped Client Site",
                client_id="client-unmapped",
                exclusion_reason="linked_client_unmapped_client_status",
            ),
            _review_row(
                UNRESOLVED_SITES_REVIEW_COLUMNS,
                source_row_number="9",
                site_id="site-dup",
                site_name_or_label="Duplicate Site A",
                exclusion_reason="duplicate_site_id_in_site_information",
            ),
            _review_row(
                UNRESOLVED_SITES_REVIEW_COLUMNS,
                source_row_number="10",
                site_id="site-dup",
                site_name_or_label="Duplicate Site B",
                exclusion_reason="duplicate_site_id_in_site_information",
            ),
        ],
    )
    _write_review_csv(
        review_dir / "client_status_review.csv",
        CLIENT_STATUS_REVIEW_COLUMNS,
        client_status_rows
        if client_status_rows is not None
        else [
            _review_row(
                CLIENT_STATUS_REVIEW_COLUMNS,
                source_row_number="4",
                client_id="client-unmapped",
                client_name_or_label="Beta LLC",
                raw_client_status="Needs Review",
            ),
        ],
    )
    _write_review_csv(
        review_dir / "duplicate_sites_review.csv",
        DUPLICATE_SITES_REVIEW_COLUMNS,
        duplicate_rows
        if duplicate_rows is not None
        else [
            _review_row(
                DUPLICATE_SITES_REVIEW_COLUMNS,
                source_row_number="9",
                site_id="site-dup",
                site_name_or_label="Duplicate Site A",
                duplicate_group="site_id:site-dup",
                keep_or_skip="skip",
            ),
            _review_row(
                DUPLICATE_SITES_REVIEW_COLUMNS,
                source_row_number="10",
                site_id="site-dup",
                site_name_or_label="Duplicate Site B",
                duplicate_group="site_id:site-dup",
                keep_or_skip="skip",
            ),
        ],
    )
    _write_review_csv(
        review_dir / "status_defaults_review.csv",
        STATUS_DEFAULTS_REVIEW_COLUMNS,
        status_default_rows
        if status_default_rows is not None
        else [
            _review_row(
                STATUS_DEFAULTS_REVIEW_COLUMNS,
                source_row_number="3",
                site_id="site-good",
                site_name_or_label="Good Site",
                raw_site_status="Needs Field Review",
                defaulted_status="active",
            ),
        ],
    )
    return review_dir


def _validation_formulas(sheet) -> set[str]:
    return {validation.formula1 for validation in sheet.data_validations.dataValidation}


def _prepare_reviewed_sample(tmp_path: Path, exports: dict[str, Path], review_dir: Path):
    return prepare_reviewed_sample(
        contacts_path=exports["contacts"],
        leads_path=exports["leads"],
        orders_path=exports["orders"],
        sites_path=exports["sites"],
        clients_template_path=TEMPLATE_DIR / "clients_template.csv",
        sites_template_path=TEMPLATE_DIR / "sites_template.csv",
        review_pack_dir=review_dir,
        clients_output_path=tmp_path / "private" / "clients_reviewed_sample.csv",
        sites_output_path=tmp_path / "private" / "sites_reviewed_sample.csv",
        report_dir=tmp_path / "reports" / "reviewed_sample",
    )


def _fake_exports(tmp_path: Path) -> dict[str, Path]:
    contacts = tmp_path / "contacts.csv"
    leads = tmp_path / "leads.csv"
    orders = tmp_path / "orders.csv"
    sites = tmp_path / "sites.csv"
    _write_monday_csv(
        contacts,
        "Contacts",
        ["Name", "First Name", "Last Name", "Email", "Phone", "Active Status", "Account", "Client ID"],
        [
            ["Alex One", "Alex", "One", "alex.one@example.com", "555-0100", "Active", "Acme Group", "client-001"],
            ["Blair Two", "Blair", "Two", "blair.two@example.com", "555-0101", "Active", "Beta LLC", "client-002"],
            ["Casey Three", "Casey", "Three", "casey.three@example.com", "555-0102", "Inactive", "Cedar Inc", "client-003"],
        ],
    )
    _write_monday_csv(
        leads,
        "Leads",
        ["Name", "Site ID", "Client ID"],
        [
            ["Lead A", "site-001", "client-001"],
            ["Lead B", "site-002", "client-001"],
            ["Lead C", "site-003", "client-002"],
            ["Lead D", "site-004", "client-003"],
            ["Lead E", "site-005", "client-003"],
            ["Lead F", "site-006", "client-003"],
        ],
    )
    _write_monday_csv(
        sites,
        "Site Information",
        ["Name", "Site ID", "Address", "CITY", "STATE", "ZIP", "Gdrive", "Status"],
        [
            ["North Basin", "site-001", "1 North Way", "Portland", "ME", "04101", "https://example.com/folders/1", "Active"],
            ["South Basin", "site-002", "2 South Way", "Portland", "ME", "04102", "", "On Hold"],
            ["West Yard", "site-003", "3 West Way", "Boston", "MA", "02108", "not-a-url", "Review Needed"],
            ["East Yard", "site-004", "4 East Way", "Dover", "NH", "03820", "", ""],
            ["Central Yard", "site-005", "5 Central Way", "Dover", "NH", "03821", "", ""],
            ["Overflow Yard", "site-006", "6 Overflow Way", "Dover", "NH", "03822", "", ""],
            ["Unlinked Yard", "site-999", "9 Missing Way", "Nowhere", "ME", "04000", "", ""],
        ],
    )
    _write_monday_csv(
        orders,
        "Order",
        ["Job Site", "SERVICE", "Scheduled date", "Job ID", "Site ID", "Client ID"],
        [["North Basin", "Inspection", "2026-01-01", "job-001", "site-001", "client-001"]],
    )
    return {"contacts": contacts, "leads": leads, "orders": orders, "sites": sites}


def _fake_review_exports(tmp_path: Path) -> dict[str, Path]:
    contacts = tmp_path / "contacts-review.csv"
    leads = tmp_path / "leads-review.csv"
    orders = tmp_path / "orders-review.csv"
    sites = tmp_path / "sites-review.csv"
    _write_monday_csv(
        contacts,
        "Contacts",
        ["Name", "First Name", "Last Name", "Email", "Phone", "Active Status", "Account", "Client ID"],
        [
            ["Alex One", "Alex", "One", "alex.one@example.com", "555-0100", "Active", "Acme Group", "client-good"],
            [
                "Blair Two",
                "Blair",
                "Two",
                "blair.two@example.com",
                "555-0101",
                "Needs Review",
                "Beta LLC",
                "client-unmapped",
            ],
            ["Casey Three", "Casey", "Three", "casey.three@example.com", "555-0102", "Active", "", "client-no-account"],
        ],
    )
    _write_monday_csv(
        leads,
        "Leads",
        ["Name", "Site ID", "Client ID"],
        [
            ["Good Lead", "site-good", "client-good"],
            ["No Client Lead", "site-no-client", ""],
            ["Unmapped Client Lead", "site-unmapped-client", "client-unmapped"],
            ["Missing Client Lead", "site-client-missing", "client-missing"],
            ["No Account Lead", "site-no-reviewable", "client-no-account"],
            ["Duplicate Lead", "site-dup", "client-good"],
        ],
    )
    _write_monday_csv(
        sites,
        "Site Information",
        ["Name", "Site ID", "Address", "CITY", "STATE", "ZIP", "Gdrive", "Status"],
        [
            ["Good Site", "site-good", "1 Good Way", "Portland", "ME", "04101", "", "Needs Field Review"],
            ["No Lead Site", "site-no-lead", "2 Missing Way", "Portland", "ME", "04102", "", "Active"],
            ["No Client Site", "site-no-client", "3 Link Way", "Portland", "ME", "04103", "", "Active"],
            ["Unmapped Client Site", "site-unmapped-client", "4 Status Way", "Boston", "MA", "02108", "", "Active"],
            ["Missing Client Site", "site-client-missing", "5 Contact Way", "Boston", "MA", "02109", "", "Active"],
            ["No Reviewable Site", "site-no-reviewable", "6 Account Way", "Dover", "NH", "03820", "", "Active"],
            ["Duplicate Site A", "site-dup", "7 Dup Way", "Dover", "NH", "03821", "", "Active"],
            ["Duplicate Site B", "site-dup", "8 Dup Way", "Dover", "NH", "03822", "", "Active"],
            ["Blank ID Site", "", "9 Blank Way", "Dover", "NH", "03823", "", "Active"],
        ],
    )
    _write_monday_csv(
        orders,
        "Order",
        ["Job Site", "SERVICE", "Scheduled date", "Job ID", "Site ID", "Client ID"],
        [["Good Site", "Inspection", "2026-01-01", "job-001", "site-good", "client-good"]],
    )
    return {"contacts": contacts, "leads": leads, "orders": orders, "sites": sites}


def test_prepare_tiny_sample_writes_valid_private_clients_and_sites(tmp_path: Path) -> None:
    exports = _fake_exports(tmp_path)
    clients_output = tmp_path / "private" / "clients_tiny_sample.csv"
    sites_output = tmp_path / "private" / "sites_tiny_sample.csv"
    report_output = tmp_path / "reports" / "monday-mapping-review.md"

    summary = prepare_tiny_sample(
        contacts_path=exports["contacts"],
        leads_path=exports["leads"],
        orders_path=exports["orders"],
        sites_path=exports["sites"],
        clients_template_path=TEMPLATE_DIR / "clients_template.csv",
        sites_template_path=TEMPLATE_DIR / "sites_template.csv",
        clients_output_path=clients_output,
        sites_output_path=sites_output,
        report_path=report_output,
        max_clients=3,
        max_sites=5,
    )

    assert summary.clients_written == 3
    assert summary.sites_written == 5
    assert summary.site_status_defaulted == 3
    clients = _read_csv(clients_output)
    sites = _read_csv(sites_output)
    client_ids = {row["client_external_id"] for row in clients}
    assert {row["source_system"] for row in clients + sites} == {"monday"}
    assert all(row["client_external_id"] in client_ids for row in sites)
    assert {row["status"] for row in sites} == {"active", "on_hold"}

    validation = validate_import_templates(clients=clients_output, sites=sites_output)
    assert validation["ready"] is True
    assert validation["totals"] == {
        "total_rows": 8,
        "valid_rows": 8,
        "invalid_rows": 0,
        "duplicate_rows": 0,
        "unresolved_references": 0,
    }

    report = report_output.read_text(encoding="utf-8")
    assert "Acme Group" not in report
    assert "North Basin" not in report
    assert "Monday Site Status Values Found" in report
    assert "Review Needed" in report
    assert "site_id_not_found_in_leads" in report
    assert "site#" in report


def test_prepare_tiny_sample_excludes_unreviewable_clients(tmp_path: Path) -> None:
    exports = _fake_exports(tmp_path)
    with exports["contacts"].open("a", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(["Dana Four", "Dana", "Four", "dana.four@example.com", "555-0104", "", "Delta Co", "client-004"])
    with exports["leads"].open("a", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(["Lead G", "site-007", "client-004"])
    with exports["sites"].open("a", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(["Review Needed", "site-007", "7 Review Way", "Salem", "MA", "01970", "", ""])

    summary = prepare_tiny_sample(
        contacts_path=exports["contacts"],
        leads_path=exports["leads"],
        orders_path=exports["orders"],
        sites_path=exports["sites"],
        clients_template_path=TEMPLATE_DIR / "clients_template.csv",
        sites_template_path=TEMPLATE_DIR / "sites_template.csv",
        clients_output_path=tmp_path / "private" / "clients.csv",
        sites_output_path=tmp_path / "private" / "sites.csv",
        report_path=tmp_path / "reports" / "review.md",
        max_clients=5,
        max_sites=10,
    )

    assert summary.client_rejections["unmapped_client_status"] == 1
    assert summary.site_rejections["linked_client_unmapped_client_status"] == 1
    assert summary.clients_written == 3
    assert summary.sites_written == 6


def test_generate_review_pack_writes_private_review_files_and_bucket_counts(tmp_path: Path) -> None:
    exports = _fake_review_exports(tmp_path)
    output_dir = tmp_path / "import_validation_reports" / "monday_review_pack"

    summary = generate_monday_review_pack(
        contacts_path=exports["contacts"],
        leads_path=exports["leads"],
        orders_path=exports["orders"],
        sites_path=exports["sites"],
        output_dir=output_dir,
        max_clients=1,
        max_sites=1,
    )

    expected_files = {
        "unresolved_sites_review.csv",
        "client_status_review.csv",
        "duplicate_sites_review.csv",
        "status_defaults_review.csv",
        "README.md",
    }
    assert {path.name for path in summary.files.values()} == expected_files
    assert all(path.exists() for path in summary.files.values())

    unresolved = _read_csv(output_dir / "unresolved_sites_review.csv")
    buckets = Counter(row["diagnostic_bucket"] for row in unresolved)
    assert buckets == {
        "Site IDs not found in Leads": 1,
        "sites with no Client ID": 1,
        "linked to clients with unmapped client status": 1,
        "Client IDs not found in Contacts": 1,
        "no reviewable client candidate": 1,
        "duplicate Site ID rows": 2,
        "blank Site ID rows": 1,
    }
    assert all("manual_client_external_id" in row for row in unresolved)
    assert all("review_status" in row for row in unresolved)

    client_status_rows = _read_csv(output_dir / "client_status_review.csv")
    assert [(row["client_id"], row["raw_client_status"], row["current_mapping"]) for row in client_status_rows] == [
        ("client-unmapped", "Needs Review", ""),
    ]

    duplicate_rows = _read_csv(output_dir / "duplicate_sites_review.csv")
    assert len(duplicate_rows) == 2
    assert {row["duplicate_group"] for row in duplicate_rows} == {"site_id:site-dup"}

    status_default_rows = _read_csv(output_dir / "status_defaults_review.csv")
    assert [(row["site_id"], row["raw_site_status"], row["defaulted_status"]) for row in status_default_rows] == [
        ("site-good", "Needs Field Review", "active"),
    ]

    readme = (output_dir / "README.md").read_text(encoding="utf-8")
    assert "Counts By Diagnostic Bucket" in readme
    assert "| Site IDs not found in Leads | 1 |" in readme
    assert "| duplicate Site ID rows | 2 |" in readme
    assert "no database writes, no provider calls" in readme


def test_generate_review_workbook_writes_expected_sheets_and_dropdowns(tmp_path: Path) -> None:
    review_dir = _write_review_pack(
        tmp_path / "reviews",
        unresolved_rows=[
            _review_row(
                UNRESOLVED_SITES_REVIEW_COLUMNS,
                source_row_number="5",
                site_id="site-no-client",
                site_name_or_label="No Client Site",
                manual_client_external_id="client-good",
                manual_status="active",
                review_status="approved",
            ),
        ],
        client_status_rows=[
            _review_row(
                CLIENT_STATUS_REVIEW_COLUMNS,
                source_row_number="4",
                client_id="client-unmapped",
                client_name_or_label="Beta LLC",
                raw_client_status="Needs Review",
                suggested_status="prospect",
                review_status="approved",
            ),
        ],
        duplicate_rows=[
            _review_row(
                DUPLICATE_SITES_REVIEW_COLUMNS,
                source_row_number="9",
                site_id="site-dup",
                site_name_or_label="Duplicate Site A",
                duplicate_group="site_id:site-dup",
                keep_or_skip="keep",
            ),
        ],
        status_default_rows=[
            _review_row(
                STATUS_DEFAULTS_REVIEW_COLUMNS,
                source_row_number="3",
                site_id="site-good",
                site_name_or_label="Good Site",
                raw_site_status="Needs Field Review",
                defaulted_status="active",
                manual_status="active",
                review_status="approved",
            ),
        ],
    )
    workbook_path = tmp_path / "reviews" / "monday_mapping_review.xlsx"

    summary = generate_monday_review_workbook(review_pack_dir=review_dir, workbook_path=workbook_path)

    assert summary.workbook_path == workbook_path
    assert summary.sheet_counts == {
        "unresolved_sites": 1,
        "client_status": 1,
        "duplicate_sites": 1,
        "status_defaults": 1,
    }
    assert summary.approved_count == 3
    assert workbook_path.exists()

    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path)
    try:
        assert workbook.sheetnames == list(REVIEW_WORKBOOK_SHEETS)

        unresolved = workbook["Unresolved Sites"]
        assert unresolved.freeze_panes == "A2"
        assert unresolved.auto_filter.ref is not None
        assert [cell.value for cell in unresolved[1]] == list(UNRESOLVED_SITES_REVIEW_COLUMNS)
        assert '"approved,skip,needs_source_fix,needs_followup"' in _validation_formulas(unresolved)
        assert '"active,inactive,on_hold,archived"' in _validation_formulas(unresolved)

        client_status = workbook["Client Status Review"]
        assert [cell.value for cell in client_status[1]] == list(CLIENT_STATUS_REVIEW_COLUMNS)
        assert '"active,inactive,prospect,archived"' in _validation_formulas(client_status)

        duplicate_sites = workbook["Duplicate Sites"]
        assert [cell.value for cell in duplicate_sites[1]] == list(DUPLICATE_SITES_REVIEW_COLUMNS)
        assert '"keep,skip,needs_source_fix,needs_followup"' in _validation_formulas(duplicate_sites)

        status_defaults = workbook["Status Defaults"]
        assert [cell.value for cell in status_defaults[1]] == list(STATUS_DEFAULTS_REVIEW_COLUMNS)
        assert '"active,inactive,on_hold,archived"' in _validation_formulas(status_defaults)

        overview = {
            row[0]: row[1]
            for row in workbook["Overview"].iter_rows(min_row=2, max_col=2, values_only=True)
            if row[0]
        }
        assert overview["unresolved sites count"] == 1
        assert overview["client status review count"] == 1
        assert overview["duplicate site review count"] == 1
        assert overview["status default review count"] == 1
        assert overview["approved count if regenerated from filled values"] == 3
        assert overview["warning"] == "This workbook is private and must not be committed."

        instructions = [
            row[1]
            for row in workbook["Instructions"].iter_rows(min_row=2, max_col=2, values_only=True)
            if row[1]
        ]
        assert "Do not import from this workbook directly." in instructions
    finally:
        workbook.close()


def test_prepare_reviewed_sample_applies_approved_mappings_and_validates(tmp_path: Path) -> None:
    exports = _fake_review_exports(tmp_path)
    review_dir = _write_review_pack(
        tmp_path / "reviews",
        unresolved_rows=[
            _review_row(
                UNRESOLVED_SITES_REVIEW_COLUMNS,
                source_row_number="5",
                site_id="site-no-client",
                site_name_or_label="No Client Site",
                manual_client_external_id="client-good",
                review_status="approved",
            ),
            _review_row(
                UNRESOLVED_SITES_REVIEW_COLUMNS,
                source_row_number="6",
                site_id="site-unmapped-client",
                site_name_or_label="Unmapped Client Site",
                manual_client_external_id="client-unmapped",
                review_status="approved",
            ),
        ],
        client_status_rows=[
            _review_row(
                CLIENT_STATUS_REVIEW_COLUMNS,
                source_row_number="4",
                client_id="client-unmapped",
                client_name_or_label="Beta LLC",
                raw_client_status="Needs Review",
                suggested_status="prospect",
                review_status="approved",
            ),
        ],
    )

    summary = _prepare_reviewed_sample(tmp_path, exports, review_dir)

    assert summary.clients_written == 2
    assert summary.sites_written == 2
    assert summary.validation_ready is True
    assert summary.validation_json_path.exists()
    assert summary.validation_md_path.exists()
    clients = _read_csv(summary.clients_path)
    sites = _read_csv(summary.sites_path)
    assert {row["client_external_id"]: row["status"] for row in clients} == {
        "client-good": "active",
        "client-unmapped": "prospect",
    }
    assert {row["site_external_id"]: row["client_external_id"] for row in sites} == {
        "site-no-client": "client-good",
        "site-unmapped-client": "client-unmapped",
    }

    validation = validate_import_templates(clients=summary.clients_path, sites=summary.sites_path)
    assert validation["ready"] is True


def test_prepare_reviewed_sample_skips_blank_review_status_rows(tmp_path: Path) -> None:
    exports = _fake_review_exports(tmp_path)
    review_dir = _write_review_pack(
        tmp_path / "reviews",
        unresolved_rows=[
            _review_row(
                UNRESOLVED_SITES_REVIEW_COLUMNS,
                source_row_number="5",
                site_id="site-no-client",
                site_name_or_label="No Client Site",
                manual_client_external_id="client-good",
                review_status="approved",
            ),
            _review_row(
                UNRESOLVED_SITES_REVIEW_COLUMNS,
                source_row_number="6",
                site_id="site-unmapped-client",
                site_name_or_label="Unmapped Client Site",
                manual_client_external_id="client-good",
                review_status="",
            ),
        ],
    )

    summary = _prepare_reviewed_sample(tmp_path, exports, review_dir)
    sites = _read_csv(summary.sites_path)

    assert summary.sites_written == 1
    assert [row["site_external_id"] for row in sites] == ["site-no-client"]


def test_prepare_reviewed_sample_skips_needs_followup_rows(tmp_path: Path) -> None:
    exports = _fake_review_exports(tmp_path)
    review_dir = _write_review_pack(
        tmp_path / "reviews",
        unresolved_rows=[
            _review_row(
                UNRESOLVED_SITES_REVIEW_COLUMNS,
                source_row_number="5",
                site_id="site-no-client",
                site_name_or_label="No Client Site",
                manual_client_external_id="client-good",
                review_status="approved",
            ),
            _review_row(
                UNRESOLVED_SITES_REVIEW_COLUMNS,
                source_row_number="6",
                site_id="site-unmapped-client",
                site_name_or_label="Unmapped Client Site",
                manual_client_external_id="client-good",
                review_status="needs_followup",
            ),
        ],
    )

    summary = _prepare_reviewed_sample(tmp_path, exports, review_dir)
    sites = _read_csv(summary.sites_path)

    assert summary.sites_written == 1
    assert [row["site_external_id"] for row in sites] == ["site-no-client"]


def test_prepare_reviewed_sample_fails_when_approved_site_lacks_manual_client(tmp_path: Path) -> None:
    exports = _fake_review_exports(tmp_path)
    review_dir = _write_review_pack(
        tmp_path / "reviews",
        unresolved_rows=[
            _review_row(
                UNRESOLVED_SITES_REVIEW_COLUMNS,
                source_row_number="5",
                site_id="site-no-client",
                site_name_or_label="No Client Site",
                review_status="approved",
            ),
        ],
    )

    with pytest.raises(ValueError, match="manual_client_external_id"):
        _prepare_reviewed_sample(tmp_path, exports, review_dir)


def test_prepare_reviewed_sample_honors_duplicate_keep_skip_decisions(tmp_path: Path) -> None:
    exports = _fake_review_exports(tmp_path)
    review_dir = _write_review_pack(
        tmp_path / "reviews",
        unresolved_rows=[
            _review_row(
                UNRESOLVED_SITES_REVIEW_COLUMNS,
                source_row_number="9",
                site_id="site-dup",
                site_name_or_label="Duplicate Site A",
                manual_client_external_id="client-good",
                review_status="approved",
            ),
            _review_row(
                UNRESOLVED_SITES_REVIEW_COLUMNS,
                source_row_number="10",
                site_id="site-dup",
                site_name_or_label="Duplicate Site B",
                manual_client_external_id="client-good",
                review_status="approved",
            ),
        ],
        duplicate_rows=[
            _review_row(
                DUPLICATE_SITES_REVIEW_COLUMNS,
                source_row_number="9",
                site_id="site-dup",
                site_name_or_label="Duplicate Site A",
                duplicate_group="site_id:site-dup",
                keep_or_skip="keep",
            ),
            _review_row(
                DUPLICATE_SITES_REVIEW_COLUMNS,
                source_row_number="10",
                site_id="site-dup",
                site_name_or_label="Duplicate Site B",
                duplicate_group="site_id:site-dup",
                keep_or_skip="skip",
            ),
        ],
    )

    summary = _prepare_reviewed_sample(tmp_path, exports, review_dir)
    sites = _read_csv(summary.sites_path)

    assert summary.sites_written == 1
    assert summary.kept_duplicate_sites == 1
    assert sites[0]["site_external_id"] == "site-dup"
    assert sites[0]["canonical_name"] == "Duplicate Site A"
    assert summary.validation_ready is True


def test_prepare_reviewed_sample_refuses_unknown_status_until_manually_mapped(tmp_path: Path) -> None:
    exports = _fake_review_exports(tmp_path)
    review_dir = _write_review_pack(
        tmp_path / "reviews",
        status_default_rows=[
            _review_row(
                STATUS_DEFAULTS_REVIEW_COLUMNS,
                source_row_number="3",
                site_id="site-good",
                site_name_or_label="Good Site",
                raw_site_status="Needs Field Review",
                defaulted_status="active",
                review_status="approved",
            ),
        ],
    )

    with pytest.raises(ValueError, match="manual_status"):
        _prepare_reviewed_sample(tmp_path, exports, review_dir)

    review_dir = _write_review_pack(
        tmp_path / "reviews-ready",
        status_default_rows=[
            _review_row(
                STATUS_DEFAULTS_REVIEW_COLUMNS,
                source_row_number="3",
                site_id="site-good",
                site_name_or_label="Good Site",
                raw_site_status="Needs Field Review",
                defaulted_status="active",
                manual_status="active",
                review_status="approved",
            ),
        ],
    )

    summary = _prepare_reviewed_sample(tmp_path, exports, review_dir)
    sites = _read_csv(summary.sites_path)

    assert summary.sites_written == 1
    assert sites[0]["site_external_id"] == "site-good"
    assert sites[0]["status"] == "active"
    assert summary.validation_ready is True


def test_prepare_reviewed_sample_does_not_open_db_or_network(tmp_path: Path, monkeypatch) -> None:
    exports = _fake_review_exports(tmp_path)
    review_dir = _write_review_pack(
        tmp_path / "reviews",
        unresolved_rows=[
            _review_row(
                UNRESOLVED_SITES_REVIEW_COLUMNS,
                source_row_number="5",
                site_id="site-no-client",
                site_name_or_label="No Client Site",
                manual_client_external_id="client-good",
                review_status="approved",
            ),
        ],
    )

    def fail_sqlite_connect(*_args, **_kwargs):
        raise AssertionError("reviewed mapping application must not open a database")

    def fail_socket_connect(*_args, **_kwargs):
        raise AssertionError("reviewed mapping application must not open a network connection")

    monkeypatch.setattr(sqlite3, "connect", fail_sqlite_connect)
    monkeypatch.setattr(socket, "create_connection", fail_socket_connect)

    summary = _prepare_reviewed_sample(tmp_path, exports, review_dir)

    assert summary.validation_ready is True
    assert list(tmp_path.rglob("*.db")) == []


def test_generate_review_pack_does_not_open_db_or_network(tmp_path: Path, monkeypatch) -> None:
    exports = _fake_review_exports(tmp_path)

    def fail_sqlite_connect(*_args, **_kwargs):
        raise AssertionError("review pack generation must not open a database")

    def fail_socket_connect(*_args, **_kwargs):
        raise AssertionError("review pack generation must not open a network connection")

    monkeypatch.setattr(sqlite3, "connect", fail_sqlite_connect)
    monkeypatch.setattr(socket, "create_connection", fail_socket_connect)

    generate_monday_review_pack(
        contacts_path=exports["contacts"],
        leads_path=exports["leads"],
        orders_path=exports["orders"],
        sites_path=exports["sites"],
        output_dir=tmp_path / "reports" / "review_pack",
        max_clients=1,
        max_sites=1,
    )

    assert list(tmp_path.rglob("*.db")) == []


def test_generate_review_workbook_does_not_open_db_or_network(tmp_path: Path, monkeypatch) -> None:
    review_dir = _write_review_pack(tmp_path / "reviews")

    def fail_sqlite_connect(*_args, **_kwargs):
        raise AssertionError("review workbook generation must not open a database")

    def fail_socket_connect(*_args, **_kwargs):
        raise AssertionError("review workbook generation must not open a network connection")

    monkeypatch.setattr(sqlite3, "connect", fail_sqlite_connect)
    monkeypatch.setattr(socket, "create_connection", fail_socket_connect)

    generate_monday_review_workbook(
        review_pack_dir=review_dir,
        workbook_path=tmp_path / "reviews" / "monday_mapping_review.xlsx",
    )

    assert list(tmp_path.rglob("*.db")) == []
