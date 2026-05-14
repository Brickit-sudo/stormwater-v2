from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urlparse


EXPECTED_EXPORTS = {
    "contacts": "Contacts_1778108333.xlsx",
    "leads": "Leads_1778108311.xlsx",
    "orders": "Order_1778108323.xlsx",
    "sites": "Site_Information_1778107854.xlsx",
}


@dataclass(frozen=True)
class TableData:
    path: Path
    headers: tuple[str, ...]
    rows: tuple[dict[str, str], ...]
    header_row: int


@dataclass(frozen=True)
class PrepSummary:
    clients_written: int
    sites_written: int
    candidate_sites: int
    site_rejections: dict[str, int]
    client_rejections: dict[str, int]
    site_status_defaulted: int
    site_urls_omitted: int
    clients_path: Path
    sites_path: Path
    report_path: Path


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value).strip()
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return re.sub(r"\s+", " ", str(value).strip())


def _normalized_name(value: Any) -> str:
    return clean(value).casefold()


def _unique_headers(raw_headers: Sequence[Any]) -> tuple[str, ...]:
    seen: dict[str, int] = {}
    headers: list[str] = []
    for index, header in enumerate(raw_headers, start=1):
        base = clean(header) or f"Column {index}"
        seen[base] = seen.get(base, 0) + 1
        headers.append(base if seen[base] == 1 else f"{base} ({seen[base]})")
    return tuple(headers)


def _detect_header_row(rows: Sequence[Sequence[Any]]) -> tuple[int, tuple[str, ...]]:
    early_rows = []
    for row_number, row in enumerate(rows[:25], start=1):
        values = [clean(value) for value in row]
        nonempty = sum(1 for value in values if value)
        early_rows.append((row_number, values, nonempty))

    candidates = [row for row in early_rows if row[2] >= 2]
    max_nonempty = max((row[2] for row in candidates), default=0)
    header_row, raw_headers = next(
        (
            (row_number, values)
            for row_number, values, nonempty in candidates
            if nonempty >= max(2, int(max_nonempty * 0.75))
        ),
        (1, early_rows[0][1] if early_rows else []),
    )
    return header_row, _unique_headers(raw_headers)


def _read_csv_table(path: Path) -> TableData:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        raw_rows = list(csv.reader(file))
    header_row, headers = _detect_header_row(raw_rows)
    rows = []
    for row_number, row in enumerate(raw_rows, start=1):
        if row_number <= header_row:
            continue
        values = [clean(value) for value in row]
        if all(value == "" for value in values):
            continue
        rows.append({"__rownum__": str(row_number), **dict(zip(headers, values[: len(headers)]))})
    return TableData(path=path, headers=headers, rows=tuple(rows), header_row=header_row)


def _read_xlsx_table(path: Path) -> TableData:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "Reading .xlsx Monday exports requires openpyxl. Export the board as CSV or run with "
            "a Python environment that has openpyxl installed.",
        ) from error

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    raw_rows = [[clean(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    workbook.close()

    header_row, headers = _detect_header_row(raw_rows)
    rows = []
    for row_number, row in enumerate(raw_rows, start=1):
        if row_number <= header_row:
            continue
        values = [clean(value) for value in row]
        if all(value == "" for value in values):
            continue
        rows.append({"__rownum__": str(row_number), **dict(zip(headers, values[: len(headers)]))})
    return TableData(path=path, headers=headers, rows=tuple(rows), header_row=header_row)


def read_table(path: Path) -> TableData:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv_table(path)
    if suffix == ".xlsx":
        return _read_xlsx_table(path)
    raise ValueError(f"Unsupported Monday export file type: {path}")


def _read_template_headers(path: Path) -> list[str]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.reader(file)
        return next(reader)


def _valid_email(value: str) -> bool:
    return bool(re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", value))


def _valid_url(value: str) -> bool:
    if not value:
        return True
    parsed = urlparse(value)
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def _map_client_status(value: str) -> str:
    text = clean(value).casefold()
    if not text:
        return ""
    if "inactive" in text:
        return "inactive"
    if "archiv" in text:
        return "archived"
    if "active" in text:
        return "active"
    if "prospect" in text or "lead" in text:
        return "prospect"
    return ""


def _first_nonempty(*values: Any) -> str:
    for value in values:
        text = clean(value)
        if text:
            return text
    return ""


def _write_csv(path: Path, fieldnames: Sequence[str], rows: Sequence[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _markdown_table(headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> list[str]:
    if not rows:
        return ["- None"]
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(str(value).replace("|", "\\|") for value in row) + " |")
    return lines


def _source_rows(tables: dict[str, TableData]) -> list[tuple[str, str, int, str]]:
    return [
        (tables["contacts"].path.name, "contacts", len(tables["contacts"].rows), "high"),
        (tables["leads"].path.name, "clients / site-client link source", len(tables["leads"].rows), "high"),
        (tables["orders"].path.name, "jobs", len(tables["orders"].rows), "high"),
        (tables["sites"].path.name, "sites", len(tables["sites"].rows), "high"),
    ]


def _render_mapping_report(
    *,
    tables: dict[str, TableData],
    client_rows: Sequence[dict[str, str]],
    site_rows: Sequence[dict[str, str]],
    candidate_sites: int,
    site_rejections: Counter[str],
    client_rejections: Counter[str],
    site_rows_by_id: dict[str, list[dict[str, str]]],
    client_account_names: dict[str, set[str]],
    site_status_defaulted: int,
    site_urls_omitted: int,
) -> str:
    missing_required = []
    if "Client ID" not in tables["contacts"].headers or "Account" not in tables["contacts"].headers:
        missing_required.append("Contacts export needs Client ID and Account for reviewed client mapping.")
    if "Site ID" not in tables["sites"].headers or "Name" not in tables["sites"].headers:
        missing_required.append("Site Information export needs Site ID and Name for reviewed site mapping.")
    if "Site ID" not in tables["leads"].headers or "Client ID" not in tables["leads"].headers:
        missing_required.append("Leads export needs Site ID and Client ID for site-client linking.")
    missing_required_lines = (
        [f"- {item}" for item in missing_required]
        if missing_required
        else ["- None for the reviewed Clients/Sites mapping path."]
    )
    duplicate_site_groups = sum(1 for rows in site_rows_by_id.values() if len(rows) > 1)
    duplicate_contact_client_groups = sum(1 for names in client_account_names.values() if len(names) > 1)

    lines = [
        "# Monday Mapping Review",
        "",
        f"- Generated at: `{datetime.now(timezone.utc).isoformat()}`",
        "- Scope: private Monday export mapping to V2 tiny Clients/Sites templates.",
        "- Privacy: no full source records, client names, site names, addresses, emails, URLs, or IDs are included here.",
        "- Safety: no database writes, no provider calls, no imports.",
        "",
        "## Files Inspected",
        *_markdown_table(["File", "Inferred Entity", "Rows", "Confidence"], _source_rows(tables)),
        "",
        "## Mapping Plan",
        *_markdown_table(
            ["V2 Target", "Source Export", "Source Columns", "Notes"],
            [
                ("clients.client_external_id / legacy_id", "Contacts", "Client ID", "Preserved from Monday source ID."),
                (
                    "clients.canonical_name",
                    "Contacts",
                    "Account",
                    "Only used when exactly one explicit Account value exists for the Client ID.",
                ),
                ("clients.status", "Contacts", "Active Status", "Mapped to V2 client statuses when possible."),
                (
                    "clients.primary_contact_name / email / phone",
                    "Contacts",
                    "First Name, Last Name, Name, Email, Phone",
                    "Optional fields; invalid email text is omitted.",
                ),
                ("sites.site_external_id / legacy_id", "Site Information", "Site ID", "Preserved from Monday source ID."),
                (
                    "sites.client_external_id",
                    "Leads",
                    "Site ID + Client ID",
                    "Only used when a Site ID maps to exactly one Client ID and that Client ID has a reviewed Account.",
                ),
                ("sites.canonical_name", "Site Information", "Name", "Required site display name."),
                ("sites.address/city/state/zip", "Site Information", "Address, CITY, STATE, ZIP", "Mapped when present."),
                ("sites.drive_folder_url", "Site Information", "Gdrive", "Only retained when it validates as http(s)."),
                (
                    "jobs",
                    "Order",
                    "Job Site, SERVICE, Scheduled date, Job ID, Site ID, Client ID",
                    "Mapping notes only; not normalized until Clients/Sites pass.",
                ),
            ],
        ),
        "",
        "## Normalized Tiny Sample",
        f"- Clients written: {len(client_rows)}",
        f"- Sites written: {len(site_rows)}",
        f"- Candidate sites ready before sample limit: {candidate_sites}",
        f"- Sites excluded or unresolved before sampling: {sum(site_rejections.values())}",
        f"- Client candidates excluded before sampling: {sum(client_rejections.values())}",
        f"- Site statuses defaulted to `active` pending Bryce review: {site_status_defaulted}",
        f"- Non-http(s) site Gdrive values omitted from selected sample: {site_urls_omitted}",
        "",
        "## Excluded / Unresolved Site Reasons",
        *_markdown_table(["Reason", "Rows Affected"], sorted(site_rejections.items())),
        "",
        "## Excluded Client Candidate Reasons",
        *_markdown_table(["Reason", "Rows Affected"], sorted(client_rejections.items())),
        "",
        "## Missing Required Columns",
        *missing_required_lines,
        "",
        "## Duplicate Candidates",
        f"- Site Information duplicate Site ID groups: {duplicate_site_groups}",
        f"- Contacts ambiguous Client ID to Account groups: {duplicate_contact_client_groups}",
        "",
        "## Columns That Need Bryce Review",
        "- Site Information `Status` did not map cleanly to V2 site statuses in this export; selected sample sites were defaulted to `active` only for validation.",
        "- Confirm whether Contacts `Account` is the right canonical V2 Client name.",
        "- Confirm whether Site Information `Gdrive` is a site folder URL when it is an http(s) URL.",
        "- Review Order export mapping after Clients/Sites pass; jobs were intentionally not normalized now.",
        "- Decide how to treat Site Information rows with no Leads `Site ID` to `Client ID` link.",
        "",
        "## Recommended Next Action",
        "- Run the tiny sample validator only. If clean, review these private CSVs manually before increasing sample size.",
        "- Do not import, do not write to the database, and do not normalize Jobs/Documents until Clients/Sites are accepted.",
    ]
    return "\n".join(lines)


def prepare_tiny_sample(
    *,
    contacts_path: Path,
    leads_path: Path,
    sites_path: Path,
    orders_path: Path,
    clients_template_path: Path,
    sites_template_path: Path,
    clients_output_path: Path,
    sites_output_path: Path,
    report_path: Path,
    max_clients: int = 5,
    max_sites: int = 10,
) -> PrepSummary:
    if max_clients < 1:
        raise ValueError("max_clients must be at least 1.")
    if max_sites < 1:
        raise ValueError("max_sites must be at least 1.")

    tables = {
        "contacts": read_table(contacts_path),
        "leads": read_table(leads_path),
        "orders": read_table(orders_path),
        "sites": read_table(sites_path),
    }

    client_account_names: dict[str, set[str]] = defaultdict(set)
    client_contact_rows: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in tables["contacts"].rows:
        client_id = clean(row.get("Client ID"))
        account = clean(row.get("Account"))
        if not client_id or not account:
            continue
        client_account_names[client_id].add(_normalized_name(account))
        client_contact_rows[client_id].append(row)

    client_candidates: dict[str, dict[str, str]] = {}
    client_rejections: Counter[str] = Counter()
    for client_id, contact_rows in client_contact_rows.items():
        account_keys = client_account_names[client_id]
        if len(account_keys) != 1:
            client_rejections["ambiguous_account_name"] += len(contact_rows)
            continue
        chosen = contact_rows[0]
        status = _map_client_status(chosen.get("Active Status", ""))
        if not status:
            client_rejections["unmapped_client_status"] += len(contact_rows)
            continue
        email = clean(chosen.get("Email"))
        if email and not _valid_email(email):
            email = ""
        contact_name = (
            " ".join(part for part in [clean(chosen.get("First Name")), clean(chosen.get("Last Name"))] if part)
            or clean(chosen.get("Name"))
        )
        client_candidates[client_id] = {
            "account": clean(chosen.get("Account")),
            "status": status,
            "primary_contact_name": contact_name,
            "email": email,
            "phone": clean(chosen.get("Phone")),
            "source_row": clean(chosen.get("__rownum__")),
        }

    site_to_client_ids: dict[str, set[str]] = defaultdict(set)
    for row in tables["leads"].rows:
        site_id = clean(row.get("Site ID"))
        client_id = clean(row.get("Client ID"))
        if site_id and client_id:
            site_to_client_ids[site_id].add(client_id)

    site_rows_by_id: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in tables["sites"].rows:
        site_id = clean(row.get("Site ID"))
        if site_id:
            site_rows_by_id[site_id].append(row)

    site_rejections: Counter[str] = Counter()
    candidate_sites: list[tuple[str, str, dict[str, str]]] = []
    for site_id, site_rows in site_rows_by_id.items():
        if len(site_rows) != 1:
            site_rejections["duplicate_site_id_in_site_information"] += len(site_rows)
            continue
        site_row = site_rows[0]
        if not clean(site_row.get("Name")):
            site_rejections["missing_site_name"] += 1
            continue
        linked_clients = site_to_client_ids.get(site_id, set())
        if len(linked_clients) != 1:
            reason = "ambiguous_site_to_client_link" if len(linked_clients) > 1 else "no_site_to_client_link"
            site_rejections[reason] += 1
            continue
        client_id = next(iter(linked_clients))
        if client_id not in client_candidates:
            site_rejections["no_reviewable_client_candidate"] += 1
            continue
        candidate_sites.append((site_id, client_id, site_row))

    sites_by_client: dict[str, list[tuple[str, str, dict[str, str]]]] = defaultdict(list)
    client_order: list[str] = []
    for site in candidate_sites:
        _site_id, client_id, _site_row = site
        sites_by_client[client_id].append(site)
        if client_id not in client_order:
            client_order.append(client_id)

    selected_clients = client_order[:max_clients]
    selected_sites: list[tuple[str, str, dict[str, str]]] = []
    per_client_first_pass = max(1, min(2, max_sites))
    for client_id in selected_clients:
        for site in sites_by_client[client_id][:per_client_first_pass]:
            if len(selected_sites) < max_sites:
                selected_sites.append(site)
    for client_id in selected_clients:
        for site in sites_by_client[client_id][per_client_first_pass:]:
            if len(selected_sites) >= max_sites:
                break
            selected_sites.append(site)
        if len(selected_sites) >= max_sites:
            break

    selected_client_ids = list(dict.fromkeys(client_id for _site_id, client_id, _site_row in selected_sites))
    selected_clients = selected_client_ids[:max_clients]
    selected_sites = [site for site in selected_sites if site[1] in set(selected_clients)][:max_sites]

    client_headers = _read_template_headers(clients_template_path)
    site_headers = _read_template_headers(sites_template_path)
    client_rows_out = []
    for client_id in selected_clients:
        client = client_candidates[client_id]
        row = {header: "" for header in client_headers}
        row.update(
            {
                "source_system": "monday",
                "legacy_source": "monday_contacts",
                "legacy_id": client_id,
                "client_external_id": client_id,
                "canonical_name": client["account"],
                "status": client["status"],
                "primary_contact_name": client["primary_contact_name"],
                "email": client["email"],
                "phone": client["phone"],
                "notes": f"Tiny validation sample only. Client account sourced from Contacts row {client['source_row']}.",
            },
        )
        client_rows_out.append(row)

    site_rows_out = []
    site_status_defaulted = 0
    site_urls_omitted = 0
    for site_id, client_id, site in selected_sites:
        drive_url = clean(site.get("Gdrive"))
        if drive_url and not _valid_url(drive_url):
            drive_url = ""
            site_urls_omitted += 1
        site_status_defaulted += 1
        notes = (
            "Tiny validation sample only. Client link proven by Leads Site ID to Client ID. "
            f"Site Information row {clean(site.get('__rownum__'))}. "
            "Status defaulted to active for validation pending Bryce review."
        )
        if not drive_url and clean(site.get("Gdrive")):
            notes += " Source Gdrive value omitted because it was not an http(s) URL."
        row = {header: "" for header in site_headers}
        row.update(
            {
                "source_system": "monday",
                "legacy_source": "monday_site_information",
                "legacy_id": site_id,
                "site_external_id": site_id,
                "client_external_id": client_id,
                "canonical_name": clean(site.get("Name")),
                "site_code": site_id,
                "status": "active",
                "address": clean(site.get("Address")),
                "city": _first_nonempty(site.get("CITY"), site.get("City")),
                "state": _first_nonempty(site.get("STATE"), site.get("State")),
                "zip": _first_nonempty(site.get("ZIP"), site.get("Zip")),
                "drive_folder_url": drive_url,
                "notes": notes,
            },
        )
        site_rows_out.append(row)

    _write_csv(clients_output_path, client_headers, client_rows_out)
    _write_csv(sites_output_path, site_headers, site_rows_out)

    report = _render_mapping_report(
        tables=tables,
        client_rows=client_rows_out,
        site_rows=site_rows_out,
        candidate_sites=len(candidate_sites),
        site_rejections=site_rejections,
        client_rejections=client_rejections,
        site_rows_by_id=site_rows_by_id,
        client_account_names=client_account_names,
        site_status_defaulted=site_status_defaulted,
        site_urls_omitted=site_urls_omitted,
    )
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(report, encoding="utf-8")

    return PrepSummary(
        clients_written=len(client_rows_out),
        sites_written=len(site_rows_out),
        candidate_sites=len(candidate_sites),
        site_rejections=dict(site_rejections),
        client_rejections=dict(client_rejections),
        site_status_defaulted=site_status_defaulted,
        site_urls_omitted=site_urls_omitted,
        clients_path=clients_output_path,
        sites_path=sites_output_path,
        report_path=report_path,
    )


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[3]


def _default_export_path(export_dir: Path, key: str) -> Path:
    return export_dir / EXPECTED_EXPORTS[key]


def _parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    repo_root = _repo_root()
    default_export_dir = repo_root / "docs" / "import_templates" / "v2" / "private" / "monday_exports"
    private_dir = repo_root / "docs" / "import_templates" / "v2" / "private"
    parser = argparse.ArgumentParser(
        description="Prepare a tiny private Monday Clients/Sites sample for V2 validation. No imports or DB writes.",
    )
    parser.add_argument("--export-dir", type=Path, default=default_export_dir)
    parser.add_argument("--contacts", type=Path)
    parser.add_argument("--leads", type=Path)
    parser.add_argument("--orders", type=Path)
    parser.add_argument("--sites", type=Path)
    parser.add_argument("--max-clients", type=int, default=5)
    parser.add_argument("--max-sites", type=int, default=10)
    parser.add_argument("--clients-output", type=Path, default=private_dir / "clients_tiny_sample.csv")
    parser.add_argument("--sites-output", type=Path, default=private_dir / "sites_tiny_sample.csv")
    parser.add_argument(
        "--report-output",
        type=Path,
        default=repo_root / "import_validation_reports" / "monday-mapping-review.md",
    )
    parser.add_argument(
        "--clients-template",
        type=Path,
        default=repo_root / "docs" / "import_templates" / "v2" / "clients_template.csv",
    )
    parser.add_argument(
        "--sites-template",
        type=Path,
        default=repo_root / "docs" / "import_templates" / "v2" / "sites_template.csv",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = _parse_args(argv)
    export_dir = args.export_dir
    contacts_path = args.contacts or _default_export_path(export_dir, "contacts")
    leads_path = args.leads or _default_export_path(export_dir, "leads")
    orders_path = args.orders or _default_export_path(export_dir, "orders")
    sites_path = args.sites or _default_export_path(export_dir, "sites")
    for path in (contacts_path, leads_path, orders_path, sites_path):
        if not path.exists():
            print(f"Missing Monday export file: {path}", file=sys.stderr)
            return 2

    summary = prepare_tiny_sample(
        contacts_path=contacts_path,
        leads_path=leads_path,
        orders_path=orders_path,
        sites_path=sites_path,
        clients_template_path=args.clients_template,
        sites_template_path=args.sites_template,
        clients_output_path=args.clients_output,
        sites_output_path=args.sites_output,
        report_path=args.report_output,
        max_clients=args.max_clients,
        max_sites=args.max_sites,
    )
    print("Prepared Monday tiny sample for validation only.")
    print(f"  Clients rows: {summary.clients_written}")
    print(f"  Sites rows:   {summary.sites_written}")
    print(f"  Clients CSV:  {summary.clients_path}")
    print(f"  Sites CSV:    {summary.sites_path}")
    print(f"  Review:       {summary.report_path}")
    print("Safety: no database writes, no imports, no provider calls.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
